"""
Builds the Sales Elevator storyline (src/engine/storylines/sales-elevator.json) from the iLead 1.0
content workbook in docs/ilead-1. Rules: docs/SIMULATION.md. Run: python3 scripts/storyline/import_ilead1.py

Money and funnel numbers are placeholders here; `npm run calibrate` sets them.
"""
import json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WB = openpyxl.load_workbook(ROOT / 'docs/ilead-1/iLead International Sales Elevator - Full content.xlsx', data_only=True)
OUT = ROOT / 'src/engine/storylines/sales-elevator.json'

STAGES = [  # Model doc phases; conversion ratios from the workbook's Phases sheet
    ('leads', 'Leads', 0.62, 'Sales Lead'), ('qualify', 'Qualify', 0.5, 'Qualify'), ('proposal', 'Proposal', 0.3, 'Proposal'),
    ('negotiation', 'Negotiation', 0.5, 'Negotiate'), ('conversion', 'Conversion', 0.5, 'Conversion')]
SHEET_STAGE = {sheet: key for key, _, _, sheet in STAGES}
TITLES = {'leads': 'Lead Generation Executive', 'qualify': 'Lead Qualifier', 'proposal': 'Proposal Specialist',
          'negotiation': 'Negotiation Specialist', 'conversion': 'Conversion Specialist'}
SHE = {'Beth Killiney', 'Rita Sandersky', 'Mandy Lobert', 'Ruth Ether', 'Sheila Frederick'}
# Authored from each actor's workbook remarks; surfaces only in conversation (SIMULATION.md 3.4).
CONCERN_LINES = {
    'kent': "Honestly? This job is not what I was promised. I expected a lot more when I joined, and nobody seems to remember that.",
    'beth': "I left Beta Elevators because of how I was managed there. So yes, I watch closely how things are run here.",
    'justin': "I have been in Lead Qualification for three years. I really want a shot at Sales Conversion.",
    'lowe': "I put a lot of technical depth into my proposals, and it never seems to get noticed.",
    'peter': "My last appraisals have not gone well. I am trying, but some encouragement would go a long way.",
}
CONCERNS = {
    'kent': 'He expected a lot more from this job, and feels it is not what he was promised.',
    'justin': 'He has been in Lead Qualification for 3 years and wants to move to Sales Conversion.',
    'peter': 'His appraisals have gone downhill and he needs encouragement, but he will not ask for it.',
    'beth': 'She left Beta Elevators over how she was managed, and watches closely how you lead.',
    'lowe': 'He feels his technical depth goes unnoticed.'
}

def clean(t):
    return re.sub(r'\s+', ' ', str(t or '')).strip()

def person_id(name):
    return clean(name).split(' ')[0].lower()

def trust_start(m, r):
    # Starting trust before calibration: follows morale, softened toward the middle.
    return max(0, min(100, round(30 + 0.4 * m + 0.15 * r)))

# Profiles
profiles = {}
for r in list(WB['Actor Profiles'].iter_rows(values_only=True))[1:]:
    if r[1]:
        profiles[clean(r[1])] = dict(previous='', tenure=clean(r[2]) or 'New joiner', experience=clean(r[3]), skills=clean(r[4]),
                                     remarks='' if clean(r[5]) in ('None.', 'None') else clean(r[5]), relations='')

# Stage values
people, group, home = {}, None, None
for r in list(WB['Actor Phase SMP'].iter_rows(values_only=True))[1:]:
    g, o, a, ph, s, m, p = r[:7]
    if g: group = clean(g)
    if o: home = clean(o)
    if not a or clean(a) == 'Actor': continue
    name = clean(a).replace('Desmond Marta', 'Desmond Mart')
    d = people.setdefault(name, {'active': group == 'Active Actors', 'home': SHEET_STAGE[home], 'byStage': {}})
    d['byStage'][SHEET_STAGE[clean(ph)]] = {'skill': int(s), 'morale': int(m), 'result': int(p)}

def to_person(name, d):
    pid = person_id(name)
    h = d['byStage'][d['home']]
    prof = profiles.get(name) or profiles.get(name.replace('Desmond Mart', 'Desmond Marta')) or dict(previous='', tenure='', experience='', skills='', remarks='', relations='')
    out = dict(id=pid, name=name, title=TITLES[d['home']], pronoun='she' if name in SHE else 'he', homeStage=d['home'],
               start=dict(h), byStage=d['byStage'], profile=prof)
    if pid in CONCERNS: out['hiddenConcern'] = CONCERNS[pid]
    if pid in CONCERN_LINES: out['concernLine'] = CONCERN_LINES[pid]
    if (ROOT / 'public' / 'assets' / 'npc' / f'{pid}.png').exists(): out['portrait'] = f'/assets/npc/{pid}.png'
    return out

members = [to_person(n, d) for n, d in people.items() if d['active']]
candidates = [to_person(n, d) for n, d in people.items() if not d['active']]

# Actions: effects per option and mismatch type, from the workbook's Actions sheet.
rows = list(WB['Actions'].iter_rows(min_row=3, values_only=True))
raw, cur, opt = {}, None, None
for r in rows:
    a, desc, o, sel, cost, restr, style, mt, s, m, p = r[:11]
    if a: cur = clean(a); raw[cur] = {'desc': clean(desc), 'cost': cost, 'options': []}; opt = None
    if o or (mt is not None and opt is None):
        opt = {'label': clean(o) if o and o != 'NULL' else raw[cur]['desc'], 'style': style, 'effects': {}}
        raw[cur]['options'].append(opt)
    if mt is not None: opt['effects'][f'm{int(mt)}'] = [int(s), int(m), int(p)]

STYLE = {1: 'D', 2: 'G', 3: 'P', 4: 'E'}
def opts(name, keys, **extra):
    res = []
    for o, k in zip(raw[name]['options'], keys):
        x = {'key': k, 'label': o['label'], 'effects': o['effects']}
        x.update(extra.get(k, {}))
        res.append(x)
    return res
def styled(name):
    return [dict(o, style=STYLE[i + 1]) for i, o in enumerate(opts(name, ['directing', 'guiding', 'partnering', 'entrusting']))]
def desc(name, fallback=''):
    return raw[name]['desc'] or fallback

actions = [
    dict(key='meet', name='Meet the team', description=desc('Meet the team'), scope='team', kind='live', format='meeting', rule='styleOption', cost=1, cooldownDays=10, options=styled('Meet the team')),
    dict(key='energize', name='Energize the team', description=desc('Energise the team'), scope='team', kind='static', rule='weeklyStyle', cost=1,
         options=opts('Energise the team', ['team_lunch', 'team_building'], team_lunch={'cooldownDays': 20}, team_building={'cooldownDays': 8})),
    dict(key='email', name='Send email', description='Write to up to 3 people. Congratulate or warn; replies arrive in your inbox.', scope='member', kind='live', format='email', rule='trend', cost=1, targets=[1, 3],
         options=opts('Send email', ['warning', 'congratulatory'], warning={'intent': 'warn'}, congratulatory={'intent': 'congratulate'})),
    dict(key='training', name='Send for training', description='Build skill away from the desk. Up to 3 people.', scope='member', kind='static', rule='training', cost=1, targets=[1, 3],
         options=opts('Send For Training', ['three_day', 'one_week'], three_day={'away': 3}, one_week={'away': 5})),
    dict(key='swap', name='Swap roles', description='Move people between stages, then explain the decision to them.', scope='member', kind='hybrid', format='roleplay', rule='swap', cost=1, targets=[1, 2], prerequisite='assess',
         options=opts('Swap / Reassign roles', ['reassign', 'swap'])),
    dict(key='hire', name='Hire member', description='Interview candidates for an open seat.', scope='team', kind='live', format='interview', rule='hire', cost=2, cooldownDays=8, unlockPeriod=3, options=opts('Hire member', ['interview'])),
    dict(key='fire', name='Let go', description='Remove someone from the team, then have the exit conversation.', scope='member', kind='hybrid', format='roleplay', rule='fire', cost=1, targets=[1, 1], options=opts('Fire member', ['let_go'])),
    dict(key='f2f', name='Meet face to face', description='A one to one conversation.', scope='member', kind='live', format='roleplay', rule='styleOption', cost=1, targets=[1, 1], options=styled('Meet Face to Face')),
    dict(key='assess', name='Assess member', description='See how this person would do in another stage. No effect on them.', scope='member', kind='static', rule='assess', cost=1, targets=[1, 1], options=opts('Assess member', ['assess'])),
    dict(key='reward', name='Reward member', description='Recognize someone, then tell them why.', scope='member', kind='hybrid', format='roleplay', rule='reward', cost=1, cooldownDays=20, targets=[1, 1], options=opts('Reward member', ['reward'])),
    dict(key='goals', name='Set goals', description='Agree goals, measures and support for the week.', scope='member', kind='live', format='plan', rule='styleOption', cost=1, targets=[1, 1], options=styled('Set Goals')),
    dict(key='coach', name='Coach member', description='Work through the job with them.', scope='member', kind='live', format='roleplay', rule='styleOption', cost=1, targets=[1, 1], options=styled('Coach member')),
    dict(key='feedback', name='Give feedback', description='Tell them how it is going and agree what changes.', scope='member', kind='live', format='chat', rule='styleOption', cost=1, targets=[1, 1], options=styled('Give feedback')),
]
for a in actions:
    if a['options'] and not a['options'][0]['label']: a['options'][0]['label'] = a['name']

# General events: 12 week schedule mapped onto 8 weeks. A week holds up to two events (days 1 and 3);
# a third moves to the next week with room.
CARD = {'Tragic accident': 'impact', 'Rumors of being acquired': 'signal', '360-degree feedback': 'signal', 'Job offer': 'capacity', 'New CRM system': 'impact',
        'Performance declines': 'diagnostic', 'Business process change': 'impact', 'Recession strikes': 'impact', 'Sales conference announcement': 'signal',
        "Strike at supplier's factories": 'impact', 'Insider trading scandal': 'signal', 'Elevator review website criticizes': 'impact'}
TARGETED = {'Job offer', 'Performance declines', 'Insider trading scandal'}
evs = []
for r in list(WB['General Events'].iter_rows(values_only=True))[1:]:
    if not r[0] or not r[3]: continue
    evs.append((int(r[3]), clean(r[0]), clean(r[1]), clean(r[2]), [int(r[5]), int(r[6]), int(r[7])]))
evs.sort()
used, events = set(), []
def tpl(t):
    return t.replace('PLACEHOLDER_ACTOR_NAME', '{name}').replace('PLACEHOLDER_CURRENT_ROLE', '{stage}')
slots = {}
for wk, name, he, she, imp in evs:
    p = max(1, round(wk * 8 / 12))
    while slots.get(p, 0) >= 2 and p < 8: p += 1
    if slots.get(p, 0) >= 2: sys.exit(f'No room for event {name}')
    sub = 1 if slots.get(p, 0) == 0 else 3
    slots[p] = slots.get(p, 0) + 1
    key = re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')
    if key[0].isdigit(): key = 'event_' + key
    events.append(dict(key=key, title=name.replace('360-degree', '360 degree'), body={'he': tpl(he), 'she': tpl(she)},
                       card=CARD[name], period=p, subPeriod=sub, impact=imp, target='member' if name in TARGETED else 'team'))

# Trigger events: messages from the workbook; conditions are engine code with these parameters.
TRIG = {'Casual leave': ('casualLeave', 1, {'resultAbove': 70, 'away': 5, 'atFractions': 0}),
        'Medical leave': ('medicalLeave', 1, {'resultAbove': 70, 'away': 2}),
        'Clueless team member': ('clueless', 1, {}),
        'Demoralized member': ('demoralized', 1, {'resultBelow': 20, 'fromFraction': 0.3}),
        'Lack of training': ('lackOfTraining', 2, {'decliningFraction': 0.5}),
        'Morale drops': ('moraleDrops', 3, {'resultAbove': 60, 'sustainedFraction': 0.25}),
        'Resignation': ('resignation', 2, {'resultBelow': 10, 'trustBelow': 15, 'moraleBelow': 20, 'fromFraction': 0.2}),
        'Role change request': ('roleChangeRequest', 3, {'sameStageFraction': 0.33, 'ignoredTrust': -4}),
        'Team member complains': ('complains', 2, {'decliningFraction': 0.25, 'trustBelow': 25})}
triggers, seen = [], set()
for r in list(WB['Trigger Events'].iter_rows(values_only=True))[1:]:
    n = clean(r[0])
    if not n or n in seen or n not in TRIG: continue
    seen.add(n)
    kind, times, params = TRIG[n]
    params = {k: v for k, v in params.items() if k != 'atFractions'}
    triggers.append(dict(kind=kind, message={'he': tpl(clean(r[1])), 'she': tpl(clean(r[2]))}, impact=[int(r[3]), int(r[4]), int(r[5])], maxTimes=times, params=params))
for r in list(WB['General Events'].iter_rows(values_only=True))[1:]:
    if clean(r[0]) == 'Training Request':
        triggers.append(dict(kind='trainingRequest', message={'he': tpl(clean(r[1])), 'she': tpl(clean(r[2]))}, maxTimes=2, params={'trust': 4}))

for a in actions:
    if a['key'] == 'swap':
        for o in a['options']:
            if o['key'] == 'reassign': o.update(targets=[1, 1], pickStage=True)
            if o['key'] == 'swap': o.update(targets=[2, 2], distinctStages=True)

storyline = dict(
    id='sales_elevator', name='Sales Elevator, Innov8 Elevators',
    money=dict(currency='USD', locale='en-US', display='symbol', target=240000, valuePerConversion=8400, inputPerSubPeriod=[8]),
    time=dict(period=dict(unit='week', count=8), costStep=1),
    stages=[dict(key=k, name=n, conversionRatio=c, ideal=2) for k, n, c, _ in STAGES],
    sponsor=dict(name='Paula Jacob', title='Regional Sales Director', styleLine='To each their own. Your people need different things from you this week.'),
    members=members, candidates=candidates,
    actions=actions,
    weeklyStyle=dict(m0=[2, 4, 5], m1=[0, -3, -4], m2=[-1, -6, -8]),
    events=events, triggers=triggers,
    performanceThreshold=10, calibrated=False)
OUT.write_text(json.dumps(storyline, indent=2, ensure_ascii=False) + '\n')
print(f'{OUT.relative_to(ROOT)}: {len(members)} members, {len(candidates)} candidates, {len(actions)} actions, {len(events)} events, {len(triggers)} triggers')
